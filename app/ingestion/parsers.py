"""Multi-format documentation parsers for Markdown, HTML, reStructuredText, PDF, Word, Excel, and PowerPoint."""

import csv
import io
import re
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from app.ingestion.chunker import ChunkDraft, MarkdownChunker
from app.ingestion.hasher import compute_sha256

# Optional third-party binary document parsers
try:
    import pypdf
except ImportError:
    pypdf = None  # type: ignore

try:
    import docx
except ImportError:
    docx = None  # type: ignore

try:
    import openpyxl
except ImportError:
    openpyxl = None  # type: ignore

try:
    import pptx
except ImportError:
    pptx = None  # type: ignore

try:
    import xlrd
except ImportError:
    xlrd = None  # type: ignore

try:
    import olefile
except ImportError:
    olefile = None  # type: ignore


def create_chunk_drafts_from_sections(
    sections: list[tuple[str, str]],
    max_words: int = 400,
    overlap_words: int = 40,
) -> list[ChunkDraft]:
    """Helper to split a list of (breadcrumb, body) tuples into semantic ChunkDraft items."""
    chunks: list[ChunkDraft] = []
    char_offset = 0

    for breadcrumb, body in sections:
        body = body.strip()
        if not body:
            continue

        words = body.split()
        if len(words) <= max_words:
            chunk_text = f"## {breadcrumb}\n\n{body}" if breadcrumb else body
            c_hash = compute_sha256(chunk_text)
            c_len = len(chunk_text)
            chunks.append(
                ChunkDraft(
                    chunk_index=len(chunks),
                    heading=breadcrumb,
                    content=chunk_text,
                    char_start=char_offset,
                    char_end=char_offset + c_len,
                    content_hash=c_hash,
                )
            )
            char_offset += c_len
        else:
            step = max_words
            for start_idx in range(0, len(words), step - overlap_words):
                sub_words = words[start_idx : start_idx + step]
                if not sub_words:
                    break
                sub_body = " ".join(sub_words)
                chunk_text = f"## {breadcrumb}\n\n{sub_body}" if breadcrumb else sub_body
                c_hash = compute_sha256(chunk_text)
                c_len = len(chunk_text)
                chunks.append(
                    ChunkDraft(
                        chunk_index=len(chunks),
                        heading=breadcrumb,
                        content=chunk_text,
                        char_start=char_offset,
                        char_end=char_offset + c_len,
                        content_hash=c_hash,
                    )
                )
                char_offset += c_len
                if start_idx + step >= len(words):
                    break

    return chunks


def _extract_strings_from_binary(data: bytes, min_len: int = 4) -> list[str]:
    """Helper to safely extract printable text runs from legacy binary files (.doc, .ppt)."""
    # Extract UTF-16LE strings
    utf16_pattern = re.compile(b"(?:[\x20-\x7e]\x00){" + str(min_len).encode() + b",}")
    utf16_matches = [m.group(0).decode("utf-16le", errors="ignore").strip() for m in utf16_pattern.finditer(data)]

    # Extract Latin-1 / ASCII printable strings
    ascii_pattern = re.compile(b"[\x20-\x7e\n\r\t]{" + str(min_len).encode() + b",}")
    ascii_matches = [m.group(0).decode("latin1", errors="ignore").strip() for m in ascii_pattern.finditer(data)]

    combined: list[str] = []
    seen: set[str] = set()
    for s in utf16_matches + ascii_matches:
        s_clean = " ".join(s.split())
        if len(s_clean) >= min_len and s_clean not in seen and not s_clean.startswith(("Microsoft", "Normal.dot")):
            seen.add(s_clean)
            combined.append(s_clean)
    return combined


class HTMLDocumentParser(HTMLParser):
    """
    Strips boilerplate and extracts structured text with hierarchical heading breadcrumbs
    from HTML documentation pages using standard library html.parser.
    """

    IGNORABLE_TAGS = {"script", "style", "nav", "footer", "header", "aside", "svg", "noscript"}
    HEADING_TAGS = {"h1": 1, "h2": 2, "h3": 3, "h4": 4, "h5": 5, "h6": 6}

    def __init__(self, max_words: int = 400):
        super().__init__()
        self.max_words = max_words
        self.ignore_depth = 0
        self.current_heading_level: int | None = None
        self.heading_buffer: list[str] = []
        self.running_headings: list[tuple[int, str]] = []  # (level, text)
        self.sections: list[tuple[str, list[str]]] = []  # (breadcrumb, paragraphs)
        self.current_paragraphs: list[str] = []
        self.in_pre = False
        self.pre_buffer: list[str] = []
        self.title: str = ""

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]):
        tag_lower = tag.lower()
        if tag_lower in self.IGNORABLE_TAGS:
            self.ignore_depth += 1
            return
        if self.ignore_depth > 0:
            return

        if tag_lower in self.HEADING_TAGS:
            self._flush_section()
            self.current_heading_level = self.HEADING_TAGS[tag_lower]
            self.heading_buffer = []
        elif tag_lower == "pre":
            self.in_pre = True
            self.pre_buffer = []

    def handle_endtag(self, tag: str):
        tag_lower = tag.lower()
        if tag_lower in self.IGNORABLE_TAGS:
            self.ignore_depth = max(0, self.ignore_depth - 1)
            return
        if self.ignore_depth > 0:
            return

        if tag_lower in self.HEADING_TAGS and self.current_heading_level is not None:
            heading_text = " ".join("".join(self.heading_buffer).split())
            if heading_text:
                if not self.title and self.current_heading_level == 1:
                    self.title = heading_text
                # Update running breadcrumb hierarchy
                while self.running_headings and self.running_headings[-1][0] >= self.current_heading_level:
                    self.running_headings.pop()
                self.running_headings.append((self.current_heading_level, heading_text))
            self.current_heading_level = None
            self.heading_buffer = []

        elif tag_lower == "pre":
            self.in_pre = False
            code_text = "".join(self.pre_buffer).strip()
            if code_text:
                self.current_paragraphs.append(f"```\n{code_text}\n```")
            self.pre_buffer = []

        elif tag_lower in ("p", "li", "tr"):
            self.current_paragraphs.append("\n")

    def handle_data(self, data: str):
        if self.ignore_depth > 0:
            return
        if self.current_heading_level is not None:
            self.heading_buffer.append(data)
        elif self.in_pre:
            self.pre_buffer.append(data)
        else:
            cleaned = data.strip()
            if cleaned:
                self.current_paragraphs.append(cleaned)

    def _flush_section(self):
        if self.current_paragraphs:
            breadcrumb = " > ".join(h[1] for h in self.running_headings) or "Document"
            combined_text = " ".join(self.current_paragraphs).strip()
            # Clean up newlines around code fences
            combined_text = re.sub(r"\s*```\s*", "\n```\n", combined_text)
            if combined_text:
                self.sections.append((breadcrumb, [combined_text]))
            self.current_paragraphs = []

    def parse(self, html_content: str, fallback_title: str = "Document") -> tuple[str, list[ChunkDraft]]:
        self.feed(html_content)
        self._flush_section()

        doc_title = self.title or fallback_title
        chunks = create_chunk_drafts_from_sections(
            [(b, "\n\n".join(ps)) for b, ps in self.sections],
            max_words=self.max_words,
        )
        return doc_title, chunks


class RSTDocumentParser:
    """
    Parses reStructuredText (.rst) documents by recognizing section underlines,
    code blocks, directives, and tables, and producing clean semantic ChunkDraft items.
    """

    ADORNMENTS = set("=-~^\"'`:#*+.")

    def __init__(self, max_words: int = 400):
        self.max_words = max_words

    def parse(self, rst_content: str, fallback_title: str = "Document") -> tuple[str, list[ChunkDraft]]:
        lines = rst_content.splitlines()
        title = fallback_title

        parsed_blocks: list[tuple[str, int, str]] = []  # (type, level, content)
        i = 0
        n = len(lines)
        levels_map: dict[str, int] = {}

        while i < n:
            line = lines[i]
            stripped = line.strip()

            # Check for heading with underline on next line
            if i + 1 < n and stripped and len(lines[i + 1].strip()) >= len(stripped):
                next_line = lines[i + 1].strip()
                first_char = next_line[0]
                if first_char in self.ADORNMENTS and all(c == first_char for c in next_line):
                    if first_char not in levels_map:
                        levels_map[first_char] = len(levels_map) + 1
                    level = levels_map[first_char]
                    if level == 1 and title == fallback_title:
                        title = stripped
                    parsed_blocks.append(("heading", level, stripped))
                    i += 2
                    continue

            # Check for code-block directive
            if stripped.startswith((".. code-block::", ".. code::")):
                lang_match = re.search(r"\.\.\s*code(?:-block)?::\s*(\w+)?", stripped)
                lang = lang_match.group(1) if lang_match and lang_match.group(1) else ""
                code_lines = []
                i += 1
                while i < n and (not lines[i].strip() or lines[i].startswith("   ") or lines[i].startswith("\t")):
                    if lines[i].strip():
                        code_lines.append(re.sub(r"^(\s{3,4}|\t)", "", lines[i]))
                    else:
                        code_lines.append("")
                    i += 1
                code_text = f"```{lang}\n" + "\n".join(code_lines).strip() + "\n```"
                parsed_blocks.append(("code", 0, code_text))
                continue

            # Check for tables (+---+ or === ===)
            if stripped.startswith("+---") or (stripped.startswith("===") and " " in stripped):
                table_lines = [line]
                i += 1
                while i < n and lines[i].strip():
                    table_lines.append(lines[i])
                    i += 1
                parsed_blocks.append(("table", 0, "\n".join(table_lines)))
                continue

            # Check for generic directives (.. note::, .. warning::)
            if stripped.startswith(".. ") and "::" in stripped:
                directive_type = stripped[3 : stripped.index("::")].strip()
                body_lines = []
                i += 1
                while i < n and (not lines[i].strip() or lines[i].startswith("   ") or lines[i].startswith("\t")):
                    if lines[i].strip():
                        body_lines.append(re.sub(r"^(\s{3,4}|\t)", "", lines[i]))
                    i += 1
                callout_text = f"> **{directive_type.upper()}**: " + " ".join(body_lines).strip()
                parsed_blocks.append(("text", 0, callout_text))
                continue

            # Regular paragraph text
            if stripped:
                para_lines = [stripped]
                i += 1
                while i < n and lines[i].strip() and not lines[i].strip().startswith(".."):
                    if i + 1 < n and lines[i + 1].strip() and lines[i + 1].strip()[0] in self.ADORNMENTS and all(c == lines[i + 1].strip()[0] for c in lines[i + 1].strip()):
                        break
                    para_lines.append(lines[i].strip())
                    i += 1
                parsed_blocks.append(("text", 0, " ".join(para_lines)))
            else:
                i += 1

        # Second pass: group parsed blocks into sections
        sections: list[tuple[str, str]] = []
        running_headings: list[tuple[int, str]] = []
        current_content: list[str] = []

        def _flush_rst():
            nonlocal current_content
            if current_content:
                breadcrumb = " > ".join(h[1] for h in running_headings) or title
                body = "\n\n".join(current_content).strip()
                if body:
                    sections.append((breadcrumb, body))
                current_content = []

        for b_type, b_level, b_text in parsed_blocks:
            if b_type == "heading":
                _flush_rst()
                while running_headings and running_headings[-1][0] >= b_level:
                    running_headings.pop()
                running_headings.append((b_level, b_text))
            else:
                current_content.append(b_text)

        _flush_rst()
        chunks = create_chunk_drafts_from_sections(sections, max_words=self.max_words)
        return title, chunks


class PDFDocumentParser:
    """
    Parses PDF documents into semantic sections and chunks with page citation anchors
    and outline bookmarks using pypdf.
    """

    def __init__(self, max_words: int = 400):
        self.max_words = max_words

    def parse(
        self,
        content: bytes | io.BytesIO | Path,
        fallback_title: str = "Document",
    ) -> tuple[str, list[ChunkDraft]]:
        if pypdf is None:
            raise ImportError("pypdf is required to parse PDF documents. Install with 'uv add pypdf'.")

        stream: io.BytesIO
        if isinstance(content, Path):
            stream = io.BytesIO(content.read_bytes())
        elif isinstance(content, bytes):
            stream = io.BytesIO(content)
        elif isinstance(content, io.BytesIO):
            stream = content
        else:
            raise TypeError(f"Unsupported content type for PDF parser: {type(content)}")

        reader = pypdf.PdfReader(stream)
        doc_title = fallback_title
        if reader.metadata and reader.metadata.title:
            cleaned_title = str(reader.metadata.title).strip()
            if cleaned_title:
                doc_title = cleaned_title

        # Extract bookmark outlines mapped to page index
        bookmarks_by_page: dict[int, list[str]] = {}

        def _extract_outline(items: list[Any]) -> None:
            for item in items:
                if isinstance(item, list):
                    _extract_outline(item)
                else:
                    try:
                        title = getattr(item, "title", str(item))
                        page_num = reader.get_destination_page_number(item)
                        if page_num is not None and page_num >= 0:
                            bookmarks_by_page.setdefault(page_num, []).append(title)
                    except Exception:
                        pass

        if reader.outline:
            try:
                _extract_outline(reader.outline)
            except Exception:
                pass

        sections: list[tuple[str, str]] = []
        running_section = ""

        for page_idx, page in enumerate(reader.pages):
            page_num = page_idx + 1
            raw_text = page.extract_text() or ""
            # Fix hyphenated words broken across line breaks
            cleaned_text = re.sub(r"(\w+)-\n(\w+)", r"\1\2", raw_text)
            cleaned_text = re.sub(r"[ \t]+", " ", cleaned_text)
            cleaned_text = re.sub(r"\n{3,}", "\n\n", cleaned_text).strip()

            if not cleaned_text:
                continue

            if page_idx in bookmarks_by_page:
                running_section = bookmarks_by_page[page_idx][0]

            if running_section:
                breadcrumb = f"{doc_title} > {running_section} (Page {page_num})"
            else:
                breadcrumb = f"{doc_title} > Page {page_num}"

            body = f"[Page {page_num}]\n\n{cleaned_text}"
            sections.append((breadcrumb, body))

        chunks = create_chunk_drafts_from_sections(sections, max_words=self.max_words)
        return doc_title, chunks


class WordDocumentParser:
    """
    Parses Microsoft Word documents (.docx and legacy .doc) into semantic sections and chunks,
    preserving heading hierarchy, paragraphs, and formatting tables as Markdown.
    """

    def __init__(self, max_words: int = 400):
        self.max_words = max_words

    def parse(
        self,
        content: bytes | io.BytesIO | Path,
        fallback_title: str = "Document",
        is_legacy_doc: bool = False,
    ) -> tuple[str, list[ChunkDraft]]:
        stream: io.BytesIO
        if isinstance(content, Path):
            is_legacy_doc = is_legacy_doc or content.suffix.lower() == ".doc"
            stream = io.BytesIO(content.read_bytes())
        elif isinstance(content, bytes):
            stream = io.BytesIO(content)
        elif isinstance(content, io.BytesIO):
            stream = content
        else:
            raise TypeError(f"Unsupported content type for Word parser: {type(content)}")

        if is_legacy_doc:
            return self._parse_legacy_doc(stream, fallback_title)
        return self._parse_docx(stream, fallback_title)

    def _parse_docx(self, stream: io.BytesIO, fallback_title: str) -> tuple[str, list[ChunkDraft]]:
        if docx is None:
            raise ImportError("python-docx is required to parse .docx documents. Install with 'uv add python-docx'.")

        doc = docx.Document(stream)
        doc_title = fallback_title
        if doc.core_properties and doc.core_properties.title:
            t = doc.core_properties.title.strip()
            if t:
                doc_title = t

        sections: list[tuple[str, list[str]]] = []
        running_headings: list[tuple[int, str]] = []
        current_paragraphs: list[str] = []

        def _flush():
            nonlocal current_paragraphs
            if current_paragraphs:
                breadcrumb = " > ".join(h[1] for h in running_headings) or doc_title
                body = "\n\n".join(current_paragraphs).strip()
                if body:
                    sections.append((breadcrumb, [body]))
                current_paragraphs = []

        # Iterate elements in document order
        for item in doc.iter_inner_content():
            if isinstance(item, docx.text.paragraph.Paragraph):
                p_text = item.text.strip()
                if not p_text:
                    continue
                style_name = (item.style.name or "").lower() if item.style else ""
                heading_level = None

                if style_name.startswith("heading"):
                    m = re.search(r"heading\s*(\d+)", style_name)
                    if m:
                        heading_level = int(m.group(1))
                elif style_name in ("title",):
                    heading_level = 1
                    if doc_title == fallback_title:
                        doc_title = p_text

                if heading_level is not None:
                    _flush()
                    while running_headings and running_headings[-1][0] >= heading_level:
                        running_headings.pop()
                    running_headings.append((heading_level, p_text))
                else:
                    current_paragraphs.append(p_text)

            elif isinstance(item, docx.table.Table):
                table_md = self._format_table(item)
                if table_md:
                    current_paragraphs.append(table_md)

        _flush()
        chunks = create_chunk_drafts_from_sections(
            [(b, "\n\n".join(ps)) for b, ps in sections],
            max_words=self.max_words,
        )
        return doc_title, chunks

    def _format_table(self, table: Any) -> str:
        rows: list[list[str]] = []
        for tr in table.rows:
            row: list[str] = []
            for tc in tr.cells:
                cell_text = " ".join(tc.text.strip().split()).replace("|", r"\|")
                row.append(cell_text)
            if any(row):
                rows.append(row)

        if not rows:
            return ""

        max_cols = max(len(r) for r in rows)
        for r in rows:
            if len(r) < max_cols:
                r.extend([""] * (max_cols - len(r)))

        lines: list[str] = [
            "| " + " | ".join(rows[0]) + " |",
            "| " + " | ".join(["---"] * max_cols) + " |",
        ]
        for r in rows[1:]:
            lines.append("| " + " | ".join(r) + " |")

        return "\n".join(lines)

    def _parse_legacy_doc(self, stream: io.BytesIO, fallback_title: str) -> tuple[str, list[ChunkDraft]]:
        raw = stream.getvalue()
        extracted = _extract_strings_from_binary(raw)
        doc_title = fallback_title
        if extracted and len(extracted[0]) < 80 and not extracted[0].startswith("http"):
            doc_title = extracted[0]

        body = "\n\n".join(extracted)
        chunks = create_chunk_drafts_from_sections(
            [(doc_title, body)],
            max_words=self.max_words,
        )
        return doc_title, chunks


class ExcelDocumentParser:
    """
    Parses spreadsheets (.xlsx, .xls, .csv, .tsv) into structured Markdown tables
    and semantic chunk drafts with repeated headers for multi-chunk tables.
    """

    def __init__(self, max_words: int = 400, rows_per_chunk: int = 40):
        self.max_words = max_words
        self.rows_per_chunk = rows_per_chunk

    def parse(
        self,
        content: bytes | io.BytesIO | Path | str,
        fallback_title: str = "Spreadsheet",
        suffix: str = ".xlsx",
    ) -> tuple[str, list[ChunkDraft]]:
        suffix = suffix.lower()
        if suffix in (".csv", ".tsv"):
            return self._parse_delimited(content, fallback_title, delimiter="," if suffix == ".csv" else "\t")

        stream: io.BytesIO
        if isinstance(content, Path):
            suffix = content.suffix.lower()
            stream = io.BytesIO(content.read_bytes())
        elif isinstance(content, bytes):
            stream = io.BytesIO(content)
        elif isinstance(content, io.BytesIO):
            stream = content
        elif isinstance(content, str):
            stream = io.BytesIO(content.encode("utf-8"))
        else:
            raise TypeError(f"Unsupported content type for Excel parser: {type(content)}")

        if suffix == ".xls":
            return self._parse_xls(stream, fallback_title)
        return self._parse_xlsx(stream, fallback_title)

    def _parse_xlsx(self, stream: io.BytesIO, fallback_title: str) -> tuple[str, list[ChunkDraft]]:
        if openpyxl is None:
            raise ImportError("openpyxl is required to parse .xlsx documents. Install with 'uv add openpyxl'.")

        wb = openpyxl.load_workbook(stream, data_only=True, read_only=True)
        doc_title = fallback_title
        if wb.properties and wb.properties.title:
            t = wb.properties.title.strip()
            if t:
                doc_title = t

        sections: list[tuple[str, str]] = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            raw_rows: list[list[str]] = []
            for row in ws.iter_rows(values_only=True):
                if row is None or not any(c is not None and str(c).strip() for c in row):
                    continue
                clean_row = [str(c).strip().replace("\n", " ").replace("|", r"\|") if c is not None else "" for c in row]
                raw_rows.append(clean_row)

            if not raw_rows:
                continue

            sheet_sections = self._rows_to_table_chunks(doc_title, sheet_name, raw_rows)
            sections.extend(sheet_sections)

        chunks = create_chunk_drafts_from_sections(sections, max_words=self.max_words)
        return doc_title, chunks

    def _parse_xls(self, stream: io.BytesIO, fallback_title: str) -> tuple[str, list[ChunkDraft]]:
        if xlrd is None:
            raise ImportError("xlrd is required to parse legacy .xls documents. Install with 'uv add xlrd'.")

        book = xlrd.open_workbook(file_contents=stream.getvalue())
        doc_title = fallback_title
        sections: list[tuple[str, str]] = []

        for sheet_idx in range(book.nsheets):
            sheet = book.sheet_by_index(sheet_idx)
            raw_rows: list[list[str]] = []
            for r in range(sheet.nrows):
                row = sheet.row_values(r)
                if not any(c is not None and str(c).strip() for c in row):
                    continue
                clean_row = [str(c).strip().replace("\n", " ").replace("|", r"\|") if c is not None else "" for c in row]
                raw_rows.append(clean_row)

            if not raw_rows:
                continue

            sheet_sections = self._rows_to_table_chunks(doc_title, sheet.name, raw_rows)
            sections.extend(sheet_sections)

        chunks = create_chunk_drafts_from_sections(sections, max_words=self.max_words)
        return doc_title, chunks

    def _parse_delimited(
        self,
        content: bytes | io.BytesIO | Path | str,
        fallback_title: str,
        delimiter: str = ",",
    ) -> tuple[str, list[ChunkDraft]]:
        text_data: str
        if isinstance(content, Path):
            text_data = content.read_text(encoding="utf-8", errors="replace")
        elif isinstance(content, bytes):
            text_data = content.decode("utf-8", errors="replace")
        elif isinstance(content, io.BytesIO):
            text_data = content.getvalue().decode("utf-8", errors="replace")
        else:
            text_data = str(content)

        reader = csv.reader(io.StringIO(text_data), delimiter=delimiter)
        raw_rows: list[list[str]] = []
        for row in reader:
            if not any(c.strip() for c in row):
                continue
            clean_row = [c.strip().replace("\n", " ").replace("|", r"\|") for c in row]
            raw_rows.append(clean_row)

        if not raw_rows:
            return fallback_title, []

        sections = self._rows_to_table_chunks(fallback_title, "Data", raw_rows)
        chunks = create_chunk_drafts_from_sections(sections, max_words=self.max_words)
        return fallback_title, chunks

    def _rows_to_table_chunks(
        self,
        doc_title: str,
        sheet_name: str,
        rows: list[list[str]],
    ) -> list[tuple[str, str]]:
        if not rows:
            return []

        header = rows[0]
        max_cols = max(len(r) for r in rows)
        if len(header) < max_cols:
            header = header + [f"Col{i+1}" for i in range(len(header), max_cols)]

        data_rows = rows[1:]
        if not data_rows:
            table_md = "| " + " | ".join(header) + " |\n| " + " | ".join(["---"] * max_cols) + " |"
            return [(f"{doc_title} > {sheet_name}", table_md)]

        sections: list[tuple[str, str]] = []
        step = self.rows_per_chunk

        for start_idx in range(0, len(data_rows), step):
            chunk_slice = data_rows[start_idx : start_idx + step]
            lines = [
                "| " + " | ".join(header) + " |",
                "| " + " | ".join(["---"] * max_cols) + " |",
            ]
            for r in chunk_slice:
                padded = r + [""] * (max_cols - len(r))
                lines.append("| " + " | ".join(padded[:max_cols]) + " |")

            table_body = "\n".join(lines)
            end_row = min(start_idx + step, len(data_rows))
            breadcrumb = f"{doc_title} > {sheet_name} (Rows {start_idx + 1}-{end_row})"
            sections.append((breadcrumb, table_body))

        return sections


class PowerPointDocumentParser:
    """
    Parses Microsoft PowerPoint presentations (.pptx and legacy .ppt) into semantic
    slide sections with titles, bullet points, tables, and speaker notes.
    """

    def __init__(self, max_words: int = 400):
        self.max_words = max_words

    def parse(
        self,
        content: bytes | io.BytesIO | Path,
        fallback_title: str = "Presentation",
        is_legacy_ppt: bool = False,
    ) -> tuple[str, list[ChunkDraft]]:
        stream: io.BytesIO
        if isinstance(content, Path):
            is_legacy_ppt = is_legacy_ppt or content.suffix.lower() == ".ppt"
            stream = io.BytesIO(content.read_bytes())
        elif isinstance(content, bytes):
            stream = io.BytesIO(content)
        elif isinstance(content, io.BytesIO):
            stream = content
        else:
            raise TypeError(f"Unsupported content type for PowerPoint parser: {type(content)}")

        if is_legacy_ppt:
            return self._parse_legacy_ppt(stream, fallback_title)
        return self._parse_pptx(stream, fallback_title)

    def _parse_pptx(self, stream: io.BytesIO, fallback_title: str) -> tuple[str, list[ChunkDraft]]:
        if pptx is None:
            raise ImportError("python-pptx is required to parse .pptx documents. Install with 'uv add python-pptx'.")

        prs = pptx.Presentation(stream)
        doc_title = fallback_title
        if prs.core_properties and prs.core_properties.title:
            t = prs.core_properties.title.strip()
            if t:
                doc_title = t

        sections: list[tuple[str, str]] = []

        for slide_idx, slide in enumerate(prs.slides, 1):
            slide_title = ""
            if slide.shapes.title and slide.shapes.title.text:
                slide_title = slide.shapes.title.text.strip()
            if not slide_title:
                slide_title = f"Slide {slide_idx}"

            body_items: list[str] = []

            for shape in slide.shapes:
                if shape == slide.shapes.title:
                    continue

                if shape.has_text_frame:
                    text_parts = []
                    for paragraph in shape.text_frame.paragraphs:
                        txt = paragraph.text.strip()
                        if txt:
                            prefix = "  * " if paragraph.level and paragraph.level > 0 else "* "
                            text_parts.append(f"{prefix}{txt}")
                    if text_parts:
                        body_items.append("\n".join(text_parts))

                if shape.has_table:
                    tbl_md = self._format_table(shape.table)
                    if tbl_md:
                        body_items.append(tbl_md)

            # Extract speaker notes if present
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()
                if notes_text:
                    body_items.append(f"> **Speaker Notes**: {notes_text}")

            content_text = "\n\n".join(body_items).strip()
            if not content_text:
                content_text = f"*(Empty Slide {slide_idx})*"

            breadcrumb = f"{doc_title} > Slide {slide_idx}: {slide_title}"
            sections.append((breadcrumb, content_text))

        chunks = create_chunk_drafts_from_sections(sections, max_words=self.max_words)
        return doc_title, chunks

    def _format_table(self, table: Any) -> str:
        rows: list[list[str]] = []
        for tr in table.rows:
            row: list[str] = []
            for tc in tr.cells:
                cell_text = " ".join(tc.text.strip().split()).replace("|", r"\|")
                row.append(cell_text)
            if any(row):
                rows.append(row)

        if not rows:
            return ""

        max_cols = max(len(r) for r in rows)
        for r in rows:
            if len(r) < max_cols:
                r.extend([""] * (max_cols - len(r)))

        lines: list[str] = [
            "| " + " | ".join(rows[0]) + " |",
            "| " + " | ".join(["---"] * max_cols) + " |",
        ]
        for r in rows[1:]:
            lines.append("| " + " | ".join(r) + " |")

        return "\n".join(lines)

    def _parse_legacy_ppt(self, stream: io.BytesIO, fallback_title: str) -> tuple[str, list[ChunkDraft]]:
        raw = stream.getvalue()
        extracted = _extract_strings_from_binary(raw)
        doc_title = fallback_title
        if extracted and len(extracted[0]) < 80:
            doc_title = extracted[0]

        body = "\n\n".join(extracted)
        chunks = create_chunk_drafts_from_sections(
            [(doc_title, body)],
            max_words=self.max_words,
        )
        return doc_title, chunks


class DocumentParserFactory:
    """Factory creating appropriate parser according to file extension."""

    @staticmethod
    def parse_file(
        file_path: Path,
        content: str | bytes | None = None,
    ) -> tuple[str, list[ChunkDraft]]:
        suffix = file_path.suffix.lower()
        fallback_title = file_path.stem.replace("_", " ").title()

        # HTML
        if suffix in (".html", ".htm"):
            text = content if isinstance(content, str) else file_path.read_text(encoding="utf-8", errors="replace")
            parser = HTMLDocumentParser()
            return parser.parse(text, fallback_title=fallback_title)

        # reStructuredText
        if suffix in (".rst", ".rest"):
            text = content if isinstance(content, str) else file_path.read_text(encoding="utf-8", errors="replace")
            parser = RSTDocumentParser()
            return parser.parse(text, fallback_title=fallback_title)

        # PDF
        if suffix == ".pdf":
            raw_bytes: bytes | io.BytesIO
            if isinstance(content, (bytes, io.BytesIO)):
                raw_bytes = content
            elif isinstance(content, str):
                raw_bytes = content.encode("latin1", errors="ignore")
            else:
                raw_bytes = file_path.read_bytes()
            pdf_parser = PDFDocumentParser()
            return pdf_parser.parse(raw_bytes, fallback_title=fallback_title)

        # Word (.docx, .doc)
        if suffix in (".docx", ".doc"):
            raw_word = content if isinstance(content, (bytes, io.BytesIO)) else file_path.read_bytes()
            word_parser = WordDocumentParser()
            return word_parser.parse(raw_word, fallback_title=fallback_title, is_legacy_doc=(suffix == ".doc"))

        # Excel (.xlsx, .xls, .csv, .tsv)
        if suffix in (".xlsx", ".xls", ".csv", ".tsv"):
            excel_parser = ExcelDocumentParser()
            data = content if content is not None else file_path
            return excel_parser.parse(data, fallback_title=fallback_title, suffix=suffix)

        # PowerPoint (.pptx, .ppt)
        if suffix in (".pptx", ".ppt"):
            raw_ppt = content if isinstance(content, (bytes, io.BytesIO)) else file_path.read_bytes()
            ppt_parser = PowerPointDocumentParser()
            return ppt_parser.parse(raw_ppt, fallback_title=fallback_title, is_legacy_ppt=(suffix == ".ppt"))

        # Default to Markdown chunker
        if isinstance(content, bytes):
            text_content = content.decode("utf-8", errors="replace")
        elif isinstance(content, str):
            text_content = content
        else:
            text_content = file_path.read_text(encoding="utf-8", errors="replace")

        chunker = MarkdownChunker()
        title = MarkdownChunker.extract_title(text_content, fallback=fallback_title)
        chunks = chunker.chunk_document(text_content, source_path=file_path.as_posix())
        return title, chunks
